import io
import time
import docx
import httpx
import pytest
from app import config, db, passwords, security, excel_export
from app.server import app

@pytest.mark.asyncio
async def test_api_routes():
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="https://testserver") as client:
        # 1. Healthz
        r = await client.get("/healthz")
        assert r.status_code == 200
        assert r.json().get("ok") is True

        # 2. Authentication setup
        username = list(config.ADMIN_CREDENTIALS.keys())[0] if config.ADMIN_CREDENTIALS else "testadmin"
        if not config.ADMIN_CREDENTIALS:
            config.ADMIN_CREDENTIALS[username] = passwords.hash_password("adminpass")
        token = security.issue_session(username)
        cookies = {"reys_session": token}
        headers = {"Origin": "https://testserver", "Referer": "https://testserver/"}

        # 3. Create report
        rep_name = f"Pytest Report {int(time.time() * 1000)}"
        r = await client.post("/api/reports", json={"name": rep_name}, cookies=cookies, headers=headers)
        assert r.status_code == 200
        rid = r.json()["report"]["id"]

        try:
            # 4. List types
            r = await client.get(f"/api/types?report_id={rid}", cookies=cookies)
            assert r.status_code == 200
            types = r.json()["types"]
            t1 = types[0]

            # 5. Submit reys with photo
            photo_bytes = b"\xff\xd8\xff\xe0\x00\x10JFIF\x00\x01\x01\x01\x00`\x00`\x00\x00\xff\xdb\x00C\x00"
            files = [("photos", ("photo.jpg", io.BytesIO(photo_bytes), "image/jpeg"))]
            data = {
                "report_id": str(rid),
                "type": t1,
                "weight": "42.0",
                "coefficient": "2.0",
                "coefficient_mode": "fixed",
                "box_weight": "2.0",
            }
            r = await client.post("/api/report", data=data, files=files, cookies=cookies, headers=headers)
            assert r.status_code == 200
            entry_id = r.json()["entry_id"]

            # 6. Fetch photo
            r = await client.get(f"/api/entry/{entry_id}/photo/0", cookies=cookies)
            assert r.status_code in (200, 307)

            # 7. Check inventory
            r = await client.get(f"/api/inventory?report_id={rid}", cookies=cookies)
            assert r.status_code == 200
            inv = r.json()["inventory"]
            assert abs(inv[t1] - 40.0) < 1e-4

            # 8. Export Excel
            for ep in ("kargo", "obshiy", "summary"):
                r = await client.get(f"/api/export/{ep}?report_id={rid}", cookies=cookies)
                assert r.status_code == 200
                assert len(r.content) > 0

            # 9. Send bulk
            r = await client.post("/api/send-bulk", json={"report_id": rid, "kind": "reys", "mode": "unsent"}, cookies=cookies, headers=headers)
            assert r.status_code == 200
            assert r.json().get("queued") >= 1

        finally:
            await client.delete(f"/api/reports/{rid}", cookies=cookies, headers=headers)


@pytest.mark.asyncio
async def test_validation_error_formatting():
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="https://testserver") as client:
        username = list(config.ADMIN_CREDENTIALS.keys())[0] if config.ADMIN_CREDENTIALS else "testadmin"
        if not config.ADMIN_CREDENTIALS:
            config.ADMIN_CREDENTIALS[username] = passwords.hash_password("adminpass")
        token = security.issue_session(username)
        cookies = {"reys_session": token}
        headers = {"Origin": "https://testserver", "Referer": "https://testserver/"}

        # Send invalid string for photos instead of file
        data = {
            "report_id": "9999",
            "type": "akb",
            "weight": "10.0",
            "coefficient": "0",
            "coefficient_mode": "none",
            "photos": "[object Object]",
        }
        r = await client.post("/api/report", data=data, cookies=cookies, headers=headers)
        assert r.status_code == 422
        res = r.json()
        assert res.get("ok") is False
        assert isinstance(res.get("detail"), str)
        assert "Rasm" in res["detail"] or "fayli" in res["detail"]


@pytest.mark.asyncio
async def test_api_report_rename():
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="https://testserver") as client:
        username = list(config.ADMIN_CREDENTIALS.keys())[0] if config.ADMIN_CREDENTIALS else "testadmin"
        if not config.ADMIN_CREDENTIALS:
            config.ADMIN_CREDENTIALS[username] = passwords.hash_password("adminpass")
        token = security.issue_session(username)
        cookies = {"reys_session": token}
        headers = {"Origin": "https://testserver", "Referer": "https://testserver/"}

        ts = int(time.time() * 1000)
        # Create 2 reports
        r1 = await client.post("/api/reports", json={"name": f"ApiRep1 {ts}"}, cookies=cookies, headers=headers)
        assert r1.status_code == 200
        rid1 = r1.json()["report"]["id"]

        r2 = await client.post("/api/reports", json={"name": f"ApiRep2 {ts}"}, cookies=cookies, headers=headers)
        assert r2.status_code == 200
        rid2 = r2.json()["report"]["id"]

        try:
            # 1. Successful rename
            patch_res = await client.patch(f"/api/reports/{rid1}", json={"name": f"ApiRep1 Renamed {ts}"}, cookies=cookies, headers=headers)
            assert patch_res.status_code == 200
            assert patch_res.json()["report"]["name"] == f"ApiRep1 Renamed {ts}"

            # 2. Duplicate name -> 409
            dup_res = await client.patch(f"/api/reports/{rid1}", json={"name": f"ApiRep2 {ts}"}, cookies=cookies, headers=headers)
            assert dup_res.status_code == 409

            # 3. Empty name -> 400
            empty_res = await client.patch(f"/api/reports/{rid1}", json={"name": "   "}, cookies=cookies, headers=headers)
            assert empty_res.status_code == 400

            # 4. Non-existent report -> 404
            nf_res = await client.patch("/api/reports/9999999", json={"name": "New Name"}, cookies=cookies, headers=headers)
            assert nf_res.status_code == 404

            # 5. Check reports list returns max >= 200
            list_res = await client.get("/api/reports", cookies=cookies)
            assert list_res.status_code == 200
            assert list_res.json()["max"] >= 200
        finally:
            await client.delete(f"/api/reports/{rid1}", cookies=cookies, headers=headers)
            await client.delete(f"/api/reports/{rid2}", cookies=cookies, headers=headers)


@pytest.mark.asyncio
async def test_cross_report_filtered_export():
    import openpyxl
    from PIL import Image

    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="https://testserver") as client:
        username = list(config.ADMIN_CREDENTIALS.keys())[0] if config.ADMIN_CREDENTIALS else "testadmin"
        if not config.ADMIN_CREDENTIALS:
            config.ADMIN_CREDENTIALS[username] = passwords.hash_password("adminpass")
        token = security.issue_session(username)
        cookies = {"reys_session": token}
        headers = {"Origin": "https://testserver", "Referer": "https://testserver/"}

        ts = int(time.time() * 1000)
        r1 = await client.post("/api/reports", json={"name": f"FilterRep1 {ts}"}, cookies=cookies, headers=headers)
        assert r1.status_code == 200
        rid1 = r1.json()["report"]["id"]

        r2 = await client.post("/api/reports", json={"name": f"FilterRep2 {ts}"}, cookies=cookies, headers=headers)
        assert r2.status_code == 200
        rid2 = r2.json()["report"]["id"]

        try:
            # Create a valid JPEG
            img = Image.new("RGB", (100, 100), color="red")
            buf = io.BytesIO()
            img.save(buf, format="JPEG")
            jpeg_bytes = buf.getvalue()

            # Submit reys in rep1
            files1 = [("photos", ("photo1.jpg", io.BytesIO(jpeg_bytes), "image/jpeg"))]
            data1 = {
                "report_id": str(rid1),
                "type": "akb",
                "weight": "120.0",
                "coefficient": "0",
                "coefficient_mode": "none",
                "box_weight": "0",
            }
            res_entry1 = await client.post("/api/report", data=data1, files=files1, cookies=cookies, headers=headers)
            assert res_entry1.status_code == 200

            # Submit reys in rep2
            files2 = [("photos", ("photo2.jpg", io.BytesIO(jpeg_bytes), "image/jpeg"))]
            data2 = {
                "report_id": str(rid2),
                "type": "akb",
                "weight": "80.0",
                "coefficient": "0",
                "coefficient_mode": "none",
                "box_weight": "0",
            }
            res_entry2 = await client.post("/api/report", data=data2, files=files2, cookies=cookies, headers=headers)
            assert res_entry2.status_code == 200

            # 1. POST /api/export/filtered (clean 4 columns without photos)
            post_res = await client.post(
                "/api/export/filtered",
                json={
                    "report_ids": [rid1, rid2],
                    "tovar_turi": "akb",
                },
                cookies=cookies,
                headers=headers,
            )
            assert post_res.status_code == 200
            assert "spreadsheetml" in post_res.headers.get("content-type", "")
            assert "AKB" in post_res.headers.get("content-disposition", "")
            assert len(post_res.content) > 0

            # Verify with openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(post_res.content))
            ws = wb.active
            assert ws.title.lower() == "akb hisoboti"
            assert ws.cell(1, 1).value == "Reys nomi"
            assert ws.cell(1, 2).value == "Og'irligi (kg)"
            assert ws.cell(1, 3).value == "Qo'shiladigan karobka (kg)"
            assert ws.cell(1, 4).value == "Jami (kg)"
            assert ws.cell(1, 5).value is None

            assert ws.cell(2, 1).value == f"FilterRep1 {ts}"
            assert float(ws.cell(2, 2).value) == 120.0
            assert ws.cell(2, 4).value == "=B2+C2"

            assert ws.cell(3, 1).value == f"FilterRep2 {ts}"
            assert float(ws.cell(3, 2).value) == 80.0
            assert ws.cell(3, 4).value == "=B3+C3"

            # Check JAMI row
            assert ws.cell(4, 1).value == "JAMI"
            assert ws.cell(4, 2).value == "=SUM(B2:B3)"
            assert ws.cell(4, 4).value == "=SUM(D2:D3)"

            # Ensure photos are NOT embedded in Excel (sent to Telegram instead)
            assert len(getattr(ws, "_images", [])) == 0

            # 2. GET /api/export/filtered
            get_res = await client.get(
                f"/api/export/filtered?report_ids={rid1},{rid2}&tovar_turi=akb",
                cookies=cookies,
                headers=headers,
            )
            assert get_res.status_code == 200
            wb_no_photo = openpyxl.load_workbook(io.BytesIO(get_res.content))
            ws_no_photo = wb_no_photo.active
            assert len(getattr(ws_no_photo, "_images", [])) == 0

            # 2b. Add Obshiy "bizda" (Bizda qoladigan) entries and verify tovar_turi == "top" takes bizda weight!
            obshiy_files1 = [("photos", ("obshiy1.jpg", io.BytesIO(jpeg_bytes), "image/jpeg"))]
            obshiy_res1 = await client.post(
                "/api/obshiy",
                data={
                    "report_id": str(rid1),
                    "section": "bizda",
                    "code": "B1",
                    "weight": "55.0",
                    "coefficient": "0",
                    "box_weight": "0",
                },
                files=obshiy_files1,
                cookies=cookies,
                headers=headers,
            )
            assert obshiy_res1.status_code == 200

            obshiy_files2 = [("photos", ("obshiy2.jpg", io.BytesIO(jpeg_bytes), "image/jpeg"))]
            obshiy_res2 = await client.post(
                "/api/obshiy",
                data={
                    "report_id": str(rid2),
                    "section": "bizda",
                    "code": "B2",
                    "weight": "45.0",
                    "coefficient": "0",
                    "box_weight": "0",
                },
                files=obshiy_files2,
                cookies=cookies,
                headers=headers,
            )
            assert obshiy_res2.status_code == 200

            top_export_res = await client.post(
                "/api/export/filtered",
                json={
                    "report_ids": [rid1, rid2],
                    "tovar_turi": "top",
                },
                cookies=cookies,
                headers=headers,
            )
            assert top_export_res.status_code == 200
            top_wb = openpyxl.load_workbook(io.BytesIO(top_export_res.content))
            top_ws = top_wb.active
            assert float(top_ws.cell(2, 2).value) == 55.0
            assert float(top_ws.cell(3, 2).value) == 45.0
            assert top_ws.cell(4, 2).value == "=SUM(B2:B3)"
            assert top_ws.cell(4, 4).value == "=SUM(D2:D3)"

            # 3. Test POST /api/send-filtered
            # 3a. Validation error: missing channel
            bad_send = await client.post(
                "/api/send-filtered",
                json={"report_ids": [rid1], "tovar_turi": "akb", "channel_id": ""},
                cookies=cookies,
                headers=headers,
            )
            assert bad_send.status_code == 400

            # 3b. Successful send with mock bot
            class MockBot:
                async def get_chat(self, chat_id):
                    return True
                async def send_photo(self, *args, **kwargs):
                    pass
                async def send_media_group(self, *args, **kwargs):
                    return []
                async def send_message(self, *args, **kwargs):
                    pass

            from app import outbox
            orig_bot = outbox._bot
            try:
                outbox.set_bot(MockBot())
                send_res = await client.post(
                    "/api/send-filtered",
                    json={
                        "report_ids": [rid1, rid2],
                        "tovar_turi": "akb",
                        "channel_id": "-1001234567890",
                    },
                    cookies=cookies,
                    headers=headers,
                )
                assert send_res.status_code == 200
                send_data = send_res.json()
                assert send_data["ok"] is True
                assert send_data["count"] == 2
                assert send_data["channel"] == "-1001234567890"

                # Verify channel was saved in app_settings & returned in /api/reports
                reps_res = await client.get("/api/reports", cookies=cookies, headers=headers)
                assert reps_res.status_code == 200
                assert reps_res.json().get("last_filter_channel") == "-1001234567890"

                # 3c. Verify Telegram isolation: x637 and xabib stay strictly separate
                await db.add_reys(rid1, "admin", "x637", 50.0, 0, 50.0, 1)
                await db.add_reys(rid1, "admin", "xabib", 20.0, 0, 20.0, 1)

                x637_send = await client.post(
                    "/api/send-filtered",
                    json={
                        "report_ids": [rid1],
                        "tovar_turi": "x637",
                        "channel_id": "-1001234567890",
                    },
                    cookies=cookies,
                    headers=headers,
                )
                assert x637_send.status_code == 200
                assert x637_send.json()["count"] == 1

                xabib_send = await client.post(
                    "/api/send-filtered",
                    json={
                        "report_ids": [rid1],
                        "tovar_turi": "xabib",
                        "channel_id": "-1001234567890",
                    },
                    cookies=cookies,
                    headers=headers,
                )
                assert xabib_send.status_code == 200
                assert xabib_send.json()["count"] == 1
            finally:
                outbox.set_bot(orig_bot)

            # 4. Validation errors for export
            bad_res1 = await client.post(
                "/api/export/filtered",
                json={"report_ids": [], "tovar_turi": "akb"},
                cookies=cookies,
                headers=headers,
            )
            assert bad_res1.status_code == 400

            bad_res2 = await client.post(
                "/api/export/filtered",
                json={"report_ids": [rid1], "tovar_turi": ""},
                cookies=cookies,
                headers=headers,
            )
            assert bad_res2.status_code == 400

        finally:
            await client.delete(f"/api/reports/{rid1}", cookies=cookies, headers=headers)
            await client.delete(f"/api/reports/{rid2}", cookies=cookies, headers=headers)


@pytest.mark.asyncio
async def test_api_adjust_kg():
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="https://testserver") as client:
        username = list(config.ADMIN_CREDENTIALS.keys())[0] if config.ADMIN_CREDENTIALS else "testadmin"
        if not config.ADMIN_CREDENTIALS:
            config.ADMIN_CREDENTIALS[username] = passwords.hash_password("adminpass")
        token = security.issue_session(username)
        cookies = {"reys_session": token}
        headers = {"Origin": "https://testserver", "Referer": "https://testserver/"}

        rep = await client.post("/api/reports", json={"name": f"KgFix Test {int(time.time() * 1000)}"}, cookies=cookies, headers=headers)
        assert rep.status_code == 200
        rid = rep.json()["report"]["id"]

        try:
            # 1. Add +6 kg to akb
            res1 = await client.post(
                f"/api/reports/{rid}/adjust-kg",
                json={"tovar_turi": "akb", "weight": "+6.0", "note": "hisobot bo'yicha farq"},
                cookies=cookies,
                headers=headers,
            )
            assert res1.status_code == 200
            data1 = res1.json()
            assert data1["ok"] is True

            # 2. Subtract -2.5 kg from akb
            res2 = await client.post(
                f"/api/reports/{rid}/adjust-kg",
                json={"tovar_turi": "akb", "weight": "-2.5", "note": "kamaytirildi"},
                cookies=cookies,
                headers=headers,
            )
            assert res2.status_code == 200
            data2 = res2.json()
            assert data2["ok"] is True

            # 2b. Verify that build_umumiy_excel applies kg_fix to Column E ("To'lashi kerak bo'lgan summa")
            content, _ = await excel_export.build_umumiy_excel(rid)
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
            ws = wb.active
            # AKB is row 5: E5 should be =C5+D5+6-2.5
            e5_val = str(ws["E5"].value)
            assert "C5+D5" in e5_val
            assert "+6" in e5_val
            assert "-2.5" in e5_val

            # 3. Verify activity record
            act_res = await client.get(f"/api/activity?report_id={rid}&start=0&end=9999999999", cookies=cookies, headers=headers)
            assert act_res.status_code == 200
            acts = act_res.json()["activity"]
            assert len(acts) == 2
            assert acts[0]["action"] == "kg_fix"
            assert acts[0]["tovar_turi"] == "AKB"
            assert acts[0]["actor"] in (username, f"pw:{username}")
            assert acts[0]["weight"] == -2.5

            # 4. Validation errors
            bad1 = await client.post(
                f"/api/reports/{rid}/adjust-kg",
                json={"tovar_turi": "akb", "weight": "0"},
                cookies=cookies,
                headers=headers,
            )
            assert bad1.status_code == 400

            bad2 = await client.post(
                f"/api/reports/{rid}/adjust-kg",
                json={"tovar_turi": "", "weight": "5"},
                cookies=cookies,
                headers=headers,
            )
            assert bad2.status_code == 400

        finally:
            await client.delete(f"/api/reports/{rid}", cookies=cookies, headers=headers)


@pytest.mark.asyncio
async def test_api_report_special_name():
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="https://testserver") as client:
        username = list(config.ADMIN_CREDENTIALS.keys())[0] if config.ADMIN_CREDENTIALS else "testadmin"
        if not config.ADMIN_CREDENTIALS:
            config.ADMIN_CREDENTIALS[username] = passwords.hash_password("adminpass")
        token = security.issue_session(username)
        cookies = {"reys_session": token}
        headers = {"Origin": "https://testserver", "Referer": "https://testserver/"}

        rep = await client.post("/api/reports", json={"name": f"SpecialName Test {int(time.time() * 1000)}"}, cookies=cookies, headers=headers)
        assert rep.status_code == 200
        rid = rep.json()["report"]["id"]

        try:
            # 1. Set special name
            res1 = await client.post(
                f"/api/reports/{rid}/special-name",
                json={"special_name": "Maxsus #1"},
                cookies=cookies,
                headers=headers,
            )
            assert res1.status_code == 200
            assert res1.json()["special_name"] == "Maxsus #1"

            # 2. Verify in list_reports
            list_res = await client.get("/api/reports", cookies=cookies)
            assert list_res.status_code == 200
            target = next((r for r in list_res.json()["reports"] if r["id"] == rid), None)
            assert target is not None
            assert target["special_name"] == "Maxsus #1"

            # 2b. Verify cell B2 in umumiy hisobot Excel export
            import openpyxl
            summary_res = await client.get(f"/api/export/summary?report_id={rid}", cookies=cookies)
            assert summary_res.status_code == 200
            wb = openpyxl.load_workbook(io.BytesIO(summary_res.content))
            ws = wb.active
            assert ws["B2"].value == "Maxsus #1"

            # 3. Clear special name
            res2 = await client.post(
                f"/api/reports/{rid}/special-name",
                json={"special_name": ""},
                cookies=cookies,
                headers=headers,
            )
            assert res2.status_code == 200
            assert res2.json()["special_name"] is None

            # 4. Validation: too long
            bad = await client.post(
                f"/api/reports/{rid}/special-name",
                json={"special_name": "x" * 65},
                cookies=cookies,
                headers=headers,
            )
            assert bad.status_code == 400

            # 5. Validation: 404 for nonexistent report
            nf = await client.post(
                "/api/reports/9999999/special-name",
                json={"special_name": "Test"},
                cookies=cookies,
                headers=headers,
            )
            assert nf.status_code == 404
        finally:
            await client.delete(f"/api/reports/{rid}", cookies=cookies, headers=headers)


@pytest.mark.asyncio
async def test_api_export_docx():
    import docx
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="https://testserver") as client:
        username = list(config.ADMIN_CREDENTIALS.keys())[0] if config.ADMIN_CREDENTIALS else "testadmin"
        if not config.ADMIN_CREDENTIALS:
            config.ADMIN_CREDENTIALS[username] = passwords.hash_password("adminpass")
        token = security.issue_session(username)
        cookies = {"reys_session": token}
        headers = {"Origin": "https://testserver", "Referer": "https://testserver/"}

        rep_name = f"Docx Test {int(time.time() * 1000)}"
        r = await client.post("/api/reports", json={"name": rep_name}, cookies=cookies, headers=headers)
        assert r.status_code == 200
        rid = r.json()["report"]["id"]

        try:
            # Set special name
            await client.post(
                f"/api/reports/{rid}/special-name",
                json={"special_name": "M184"},
                cookies=cookies,
                headers=headers,
            )

            # Export docx
            res = await client.get(f"/api/export/docx?report_id={rid}", cookies=cookies)
            assert res.status_code == 200
            assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in res.headers.get("content-type", "")
            assert "attachment;" in res.headers.get("content-disposition", "")
            assert ".docx" in res.headers.get("content-disposition", "")

            # Verify document can be parsed
            doc = docx.Document(io.BytesIO(res.content))
            assert len(doc.paragraphs) >= 4
            assert "AVIA M184" in doc.paragraphs[0].text
            assert "TOP CARGO M184" in doc.paragraphs[1].text
            assert "Avia M184" in doc.paragraphs[2].text
            assert "(O'zimizga qolgan.)" in doc.paragraphs[3].text
        finally:
            await client.delete(f"/api/reports/{rid}", cookies=cookies, headers=headers)


@pytest.mark.asyncio
async def test_custom_types_umumiy_excel():
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="https://testserver") as client:
        username = list(config.ADMIN_CREDENTIALS.keys())[0] if config.ADMIN_CREDENTIALS else "testadmin"
        if not config.ADMIN_CREDENTIALS:
            config.ADMIN_CREDENTIALS[username] = passwords.hash_password("adminpass")
        token = security.issue_session(username)
        cookies = {"reys_session": token}
        headers = {"Origin": "https://testserver", "Referer": "https://testserver/"}

        rep_name = f"Custom Cargo Test {int(time.time() * 1000)}"
        r = await client.post("/api/reports", json={"name": rep_name}, cookies=cookies, headers=headers)
        assert r.status_code == 200
        rid = r.json()["report"]["id"]

        try:
            # 1. Add standard cargo (AKB)
            await db.add_reys(rid, "tester", "AKB", 50.0, 1.0, 49.0, 0)

            # 2. Add custom cargo (XON CARGO)
            await db.add_reys(rid, "tester", "XON CARGO", 10.0, 1.0, 9.0, 0)

            # 3. Add bizda entry in obshiy ves (actor="tester", action="bizda")
            await db.add_obshiy(rid, "tester", "bizda", "001", 100.0, 0, 100.0, 0, 2.0)

            # 4. Generate umumiy excel
            content, fname = await excel_export.build_umumiy_excel(rid)
            assert content is not None
            assert "UMUMIY HISOBOT" in fname

            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
            ws = wb.active

            # Find row for XON CARGO
            xon_row = None
            for row in range(4, ws.max_row + 1):
                val = str(ws.cell(row, 2).value or "").strip().upper()
                if "XON" in val:
                    xon_row = row
                    break
            assert xon_row is not None
            assert ws.cell(xon_row, 3).value == 9.0
            assert ws.cell(xon_row, 4).value == f"=C{xon_row}*$H$3"
            assert ws.cell(xon_row, 5).value == f"=C{xon_row}+D{xon_row}"

            # Check Mandarin C4 formula includes -C{xon_row}
            c4_formula = str(ws["C4"].value)
            assert f"C{xon_row}" in c4_formula

            # Check G3 formula does NOT include C{xon_row}
            g3_formula = str(ws["G3"].value)
            assert f"C{xon_row}" not in g3_formula

            # 5. Check calculate_report_metrics
            metrics = await excel_export.calculate_report_metrics(rid)
            assert "xon cargo" in metrics
            assert metrics["xon cargo"]["weight"] == 9.0
            assert metrics["xon cargo"]["box_weight"] > 0
            assert metrics["xon cargo"]["total"] > 9.0

            # 6. Check build_special_docx uses Column E (total)
            docx_bytes, docx_name = await excel_export.build_special_docx(rid)
            doc = docx.Document(io.BytesIO(docx_bytes))
            doc_text = "\n".join(p.text for p in doc.paragraphs)
            assert "XON CARGO" in doc_text
            xon_tot_str = f"{round(metrics['xon cargo']['total'], 2):.2f} KG"
            assert xon_tot_str in doc_text
            assert "XON CARGO - 9.00 KG" not in doc_text
        finally:
            await client.delete(f"/api/reports/{rid}", cookies=cookies, headers=headers)





